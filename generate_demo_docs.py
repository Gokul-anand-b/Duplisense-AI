import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib import colors

output_dir = r"c:\Users\bgoku\OneDrive\Desktop\final year project\demo_documents"
os.makedirs(output_dir, exist_ok=True)

# =============================================================================
# 1. CREATE PDF: Enterprise_RAG_Document_Pipeline.pdf
# =============================================================================
pdf_path = os.path.join(output_dir, "Enterprise_RAG_Document_Pipeline.pdf")
doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
styles = getSampleStyleSheet()

title_style = ParagraphStyle(
    'DocTitle',
    parent=styles['Heading1'],
    fontSize=18,
    leading=22,
    textColor=colors.HexColor('#1e1b4b'),
    spaceAfter=12
)
heading_style = ParagraphStyle(
    'DocHeading',
    parent=styles['Heading2'],
    fontSize=12,
    leading=16,
    textColor=colors.HexColor('#4338ca'),
    spaceBefore=10,
    spaceAfter=4
)
body_style = ParagraphStyle(
    'DocBody',
    parent=styles['Normal'],
    fontSize=10,
    leading=14,
    textColor=colors.HexColor('#334155'),
    spaceAfter=6
)

story = [
    Paragraph("Project Title: Enterprise Semantic Document Parser & RAG Knowledge Pipeline", title_style),
    Paragraph("Organization: Cloud Infrastructure & Platform Engineering Group", body_style),
    Paragraph("Date: March 2026 | Status: Architecture Review", body_style),
    Spacer(1, 10),
    
    Paragraph("Description:", heading_style),
    Paragraph("A high-throughput distributed microservice architecture designed for automated parsing of unstructured PDF and DOCX technical reports. The platform incorporates sentence-transformers for 768-dimensional dense vector embeddings, FAISS vector indexing for similarity queries, and Celery asynchronous queues for high-concurrency background extraction.", body_style),
    
    Paragraph("Problem Statement:", heading_style),
    Paragraph("Multiple internal product teams repeatedly build redundant document parsing logic and isolated vector search indexes, resulting in severe resource fragmentation, duplicate cloud costs, and high ingestion latency during quarterly vendor document intake.", body_style),
    
    Paragraph("Core Objectives:", heading_style),
    Paragraph("1. Deliver standardized document chunking and metadata extraction across all corporate formats.<br/>2. Accelerate semantic vector retrieval latency to under 15 milliseconds via persistent FAISS indexing.<br/>3. Establish unified REST and gRPC endpoints for cross-team service reuse.", body_style),
    
    Paragraph("Technology Stack:", heading_style),
    Paragraph("Programming Languages: Python, TypeScript, SQL<br/>Frameworks & Libraries: FastAPI, React, Celery, LangChain<br/>Databases & Storage: PostgreSQL, Redis, pgvector, FAISS<br/>APIs & Cloud: OpenAI API, AWS S3, REST API, Docker, Kubernetes<br/>AI / ML Tools: Sentence-Transformers, PyTorch, Vector Embeddings, RAG", body_style),
    Spacer(1, 6),
    
    Paragraph("Repository & Documentation:", heading_style),
    Paragraph("GitHub Repository: https://github.com/enterprise-org/rag-doc-pipeline<br/>Documentation: https://docs.internal/architecture/rag-pipeline", body_style),
]

doc.build(story)
print("Created PDF:", pdf_path)

# =============================================================================
# 2. CREATE DOCX: Smart_Healthcare_Telemetry_Hub.docx
# =============================================================================
docx_path = os.path.join(output_dir, "Smart_Healthcare_Telemetry_Hub.docx")
docx = Document()

docx.add_heading("Smart Healthcare Telemetry & Remote Patient Monitoring Hub", level=1)
docx.add_paragraph("Department of Medical Informatics | Systems Engineering")

p_desc_h = docx.add_heading("Description:", level=2)
docx.add_paragraph(
    "A real-time edge-to-cloud IoT telemedicine platform for continuous ingestion, automated anomaly detection, "
    "and visualization of biometric vital signs from wearable medical sensors. Built with high-availability microservices "
    "and automated clinical escalation workflows."
)

p_prob_h = docx.add_heading("Problem Statement:", level=2)
docx.add_paragraph(
    "Critical care units face clinical alert fatigue and delayed intervention times due to siloed telemetry systems "
    "that fail to synthesize multiparametric physiological data in real time."
)

p_obj_h = docx.add_heading("Core Objectives:", level=2)
docx.add_paragraph("1. Stream biometric events with sub-50ms latency across 5,000 concurrent patient feeds.", style='List Bullet')
docx.add_paragraph("2. Implement automated anomaly scoring using lightweight PyTorch temporal models.", style='List Bullet')
docx.add_paragraph("3. Provide FHIR-compliant REST and WebSocket endpoints for hospital EHR integration.", style='List Bullet')

p_tech_h = docx.add_heading("Technology Stack:", level=2)
docx.add_paragraph("Programming Languages: Python, Go, TypeScript")
docx.add_paragraph("Frameworks & Libraries: FastAPI, React 19, Celery")
docx.add_paragraph("Database & Storage: PostgreSQL, Redis, TimescaleDB, MongoDB")
docx.add_paragraph("APIs & Infrastructure: WebSocket, Docker, Kubernetes, REST API")
docx.add_paragraph("AI/ML Components: PyTorch, scikit-learn, Vector Embeddings")

docx.add_heading("Repository & Documentation:", level=2)
docx.add_paragraph("GitHub: https://github.com/health-systems/telemetry-hub")
docx.add_paragraph("Documentation: https://docs.hospital-health.internal/telemetry")

docx.save(docx_path)
print("Created DOCX:", docx_path)

# =============================================================================
# 3. CREATE TXT: Autonomous_Robotics_Perception_System.txt
# =============================================================================
txt_path = os.path.join(output_dir, "Autonomous_Robotics_Perception_System.txt")
content_txt = """Autonomous Mobile Robotics Spatial Perception and Mapping Architecture
Advanced Robotics and Autonomous Systems Laboratory

Abstract:
A unified spatial perception and point-cloud feature extraction pipeline for autonomous mobile robots operating in dynamic warehouse environments. The system combines LiDAR odometry, stereo depth estimation, and semantic obstacle classification to maintain a global topological vector map.

Problem Definition:
Conventional SLAM architectures experience rapid localization drift and severe computational bottlenecks when navigating unstructured logistics hubs with moving inventory and variable lighting.

Key Deliverables:
1. Real-time 3D bounding box detection running at 30 FPS on embedded edge hardware.
2. Maintain sub-centimeter global localization fidelity over 10-kilometer mapping trajectories.
3. Expose zero-copy shared memory IPC and gRPC APIs for fleet management software.

Technical Specifications:
Languages: C++, Python, CUDA
Frameworks: PyTorch, OpenCV, ROS2
Databases & Cache: Redis, SQLite
APIs & Connectivity: gRPC, Docker, WebSocket
AI & ML Engine: YOLO, Detectron2, Vector Embeddings

Source Code: https://github.com/robotics-lab/spatial-perception
System Docs: https://docs.robotics.internal/spatial-mapping/v1
"""

with open(txt_path, "w", encoding="utf-8") as f:
    f.write(content_txt)

print("Created TXT:", txt_path)
print("All 3 demo documents created successfully in demo_documents/!")
