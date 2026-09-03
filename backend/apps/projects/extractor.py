"""
DupliSense AI — Document Extraction Engine
============================================
3-Layer extraction pipeline using:
  1. Text Extraction (pdfplumber / python-docx / txt)
  2. Section-Aware Regex Parser (Rule-Based Information Extraction - RBIE)
  3. Heuristic Keyword Matcher (Aho-Corasick Automaton + TF-IDF Paragraph Scorer)

Extracts: title, description, problem_statement, objectives,
          programming_languages, frameworks, database_tech,
          apis_used, ai_ml_tech, github_url, documentation_url
"""

import re
import os
import math
from collections import defaultdict, deque


# =============================================================================
# TECHNOLOGY DICTIONARY — Curated list of 200+ technologies with aliases
# =============================================================================

TECH_DICTIONARY = {
    "programming_languages": {
        "Python": ["python", "python3", "cpython", "py3"],
        "JavaScript": ["javascript", "ecmascript", "es6", "es2015"],
        "TypeScript": ["typescript"],
        "Java": ["java", "openjdk", "jdk"],
        "C++": ["c++", "cpp", "cplusplus"],
        "C#": ["c#", "csharp", "c sharp"],
        "Go": ["golang"],
        "Rust": ["rust", "rustlang"],
        "Ruby": ["ruby"],
        "PHP": ["php"],
        "Swift": ["swift"],
        "Kotlin": ["kotlin"],
        "Scala": ["scala"],
        "R": ["r language", "r programming", "rlang"],
        "Dart": ["dart"],
        "Lua": ["lua"],
        "Perl": ["perl"],
        "Haskell": ["haskell"],
        "Elixir": ["elixir"],
        "Clojure": ["clojure"],
        "Julia": ["julia"],
        "MATLAB": ["matlab"],
        "SQL": ["sql", "structured query language"],
        "Shell": ["bash", "shell script", "zsh"],
        "Objective-C": ["objective-c", "objc"],
        "Assembly": ["assembly", "asm"],
        "Groovy": ["groovy"],
        "F#": ["f#", "fsharp"],
        "Erlang": ["erlang"],
        "Zig": ["zig"],
        "Solidity": ["solidity"],
        "VHDL": ["vhdl"],
        "Fortran": ["fortran"],
        "COBOL": ["cobol"],
        "Prolog": ["prolog"],
        "Lisp": ["lisp", "common lisp"],
    },
    "frameworks": {
        "React": ["react", "react.js", "reactjs", "react 18", "react 19"],
        "Next.js": ["next.js", "nextjs"],
        "Vue.js": ["vue.js", "vuejs", "vue 3"],
        "Angular": ["angular", "angularjs"],
        "Svelte": ["svelte", "sveltekit"],
        "Django": ["django"],
        "Django REST Framework": ["django rest framework", "drf", "django rest"],
        "Flask": ["flask"],
        "FastAPI": ["fastapi", "fast api"],
        "Express.js": ["express.js", "expressjs", "express"],
        "NestJS": ["nestjs", "nest.js"],
        "Spring Boot": ["spring boot", "springboot"],
        "Spring": ["spring framework"],
        "Ruby on Rails": ["ruby on rails", "rails", "ror"],
        "Laravel": ["laravel"],
        "ASP.NET": ["asp.net", "aspnet", "asp.net core"],
        ".NET": [".net", "dotnet", ".net core"],
        "Gin": ["gin framework", "gin-gonic"],
        "Fiber": ["fiber", "gofiber"],
        "Echo": ["echo framework"],
        "Actix": ["actix", "actix-web"],
        "Rocket": ["rocket framework"],
        "Phoenix": ["phoenix framework", "phoenix elixir"],
        "Remix": ["remix"],
        "Nuxt.js": ["nuxt.js", "nuxtjs", "nuxt"],
        "Gatsby": ["gatsby"],
        "Astro": ["astro"],
        "Flutter": ["flutter"],
        "React Native": ["react native"],
        "Ionic": ["ionic"],
        "Electron": ["electron"],
        "Tauri": ["tauri"],
        "Qt": ["qt framework"],
        "Celery": ["celery"],
        "LangChain": ["langchain", "lang chain"],
        "LlamaIndex": ["llamaindex", "llama index"],
        "Streamlit": ["streamlit"],
        "Gradio": ["gradio"],
        "Bootstrap": ["bootstrap"],
        "Tailwind CSS": ["tailwindcss", "tailwind css", "tailwind"],
        "Material UI": ["material ui", "mui", "material-ui"],
        "Chakra UI": ["chakra ui", "chakra-ui"],
        "Ant Design": ["ant design", "antd"],
        "Vite": ["vite"],
        "Webpack": ["webpack"],
        "Storybook": ["storybook"],
        "Jest": ["jest"],
        "Pytest": ["pytest"],
        "Mocha": ["mocha"],
        "Cypress": ["cypress"],
        "Playwright": ["playwright"],
        "Selenium": ["selenium"],
    },
    "database_tech": {
        "PostgreSQL": ["postgresql", "postgres"],
        "MySQL": ["mysql"],
        "MariaDB": ["mariadb"],
        "SQLite": ["sqlite", "sqlite3"],
        "Oracle DB": ["oracle db", "oracle database", "oracledb"],
        "SQL Server": ["sql server", "mssql", "microsoft sql"],
        "MongoDB": ["mongodb", "mongo"],
        "Redis": ["redis"],
        "Cassandra": ["cassandra"],
        "DynamoDB": ["dynamodb"],
        "CouchDB": ["couchdb"],
        "Neo4j": ["neo4j"],
        "ArangoDB": ["arangodb"],
        "InfluxDB": ["influxdb"],
        "TimescaleDB": ["timescaledb"],
        "Elasticsearch": ["elasticsearch", "elastic search"],
        "OpenSearch": ["opensearch"],
        "Solr": ["solr", "apache solr"],
        "Firebase Realtime DB": ["firebase realtime", "firebase rtdb"],
        "Firestore": ["firestore", "cloud firestore"],
        "Supabase": ["supabase"],
        "PlanetScale": ["planetscale"],
        "CockroachDB": ["cockroachdb"],
        "Snowflake": ["snowflake"],
        "BigQuery": ["bigquery", "big query"],
        "Redshift": ["redshift", "amazon redshift"],
        "pgvector": ["pgvector", "pg_vector"],
        "Pinecone": ["pinecone"],
        "Weaviate": ["weaviate"],
        "Qdrant": ["qdrant"],
        "Chroma": ["chroma", "chromadb"],
        "Milvus": ["milvus"],
        "Redis Cluster": ["redis cluster"],
        "Memcached": ["memcached"],
        "RabbitMQ": ["rabbitmq"],
        "Apache Kafka": ["kafka", "apache kafka"],
        "Amazon S3": ["amazon s3", "aws s3", "s3 bucket"],
        "MinIO": ["minio"],
        "HDFS": ["hdfs", "hadoop distributed"],
    },
    "apis_used": {
        "OpenAI API": ["openai api", "openai", "gpt-4", "gpt-3.5", "chatgpt api", "gpt-4o", "gpt api"],
        "Anthropic API": ["anthropic", "claude api", "claude"],
        "Google Gemini API": ["gemini api", "google gemini", "gemini"],
        "Hugging Face API": ["hugging face", "huggingface", "hf api"],
        "Cohere API": ["cohere"],
        "AWS Services": ["aws", "amazon web services"],
        "AWS Lambda": ["aws lambda", "lambda function"],
        "AWS SQS": ["aws sqs", "sqs"],
        "AWS SNS": ["aws sns", "sns"],
        "Google Cloud Platform": ["gcp", "google cloud"],
        "Azure": ["azure", "microsoft azure"],
        "Stripe API": ["stripe"],
        "PayPal API": ["paypal"],
        "Razorpay API": ["razorpay"],
        "Twilio API": ["twilio"],
        "SendGrid": ["sendgrid"],
        "Mailgun": ["mailgun"],
        "Firebase Auth": ["firebase auth", "firebase authentication"],
        "Auth0": ["auth0"],
        "Okta": ["okta"],
        "OAuth2": ["oauth2", "oauth 2.0", "oauth"],
        "JWT": ["jwt", "json web token"],
        "REST API": ["rest api", "restful api", "restful"],
        "GraphQL": ["graphql", "graph ql"],
        "gRPC": ["grpc", "g-rpc"],
        "WebSocket": ["websocket", "web socket"],
        "Socket.IO": ["socket.io", "socketio"],
        "Docker": ["docker"],
        "Kubernetes": ["kubernetes", "k8s"],
        "Nginx": ["nginx"],
        "Apache": ["apache http", "apache server"],
        "GitHub API": ["github api"],
        "GitLab API": ["gitlab api"],
        "Slack API": ["slack api", "slack"],
        "Discord API": ["discord api", "discord bot"],
        "Telegram Bot API": ["telegram bot", "telegram api"],
        "Google Maps API": ["google maps", "maps api"],
        "Mapbox": ["mapbox"],
        "Cloudinary": ["cloudinary"],
        "Imgur API": ["imgur"],
        "YouTube API": ["youtube api", "youtube data api"],
        "Twitter API": ["twitter api", "x api"],
        "LinkedIn API": ["linkedin api"],
        "Elasticsearch API": ["elasticsearch api"],
    },
    "ai_ml_tech": {
        "TensorFlow": ["tensorflow", "tf2", "tf 2.0"],
        "PyTorch": ["pytorch", "torch"],
        "Keras": ["keras"],
        "scikit-learn": ["scikit-learn", "sklearn", "scikit learn"],
        "XGBoost": ["xgboost"],
        "LightGBM": ["lightgbm"],
        "CatBoost": ["catboost"],
        "Hugging Face Transformers": ["hugging face transformers", "hf transformers", "transformers library"],
        "Sentence-Transformers": ["sentence-transformers", "sentence transformers", "sbert"],
        "FAISS": ["faiss", "facebook ai similarity search"],
        "SpaCy": ["spacy"],
        "NLTK": ["nltk", "natural language toolkit"],
        "Gensim": ["gensim"],
        "OpenCV": ["opencv", "cv2"],
        "Pillow": ["pillow", "pil"],
        "MediaPipe": ["mediapipe"],
        "YOLO": ["yolo", "yolov5", "yolov8", "ultralytics"],
        "Detectron2": ["detectron2"],
        "Stable Diffusion": ["stable diffusion"],
        "DALL-E": ["dall-e", "dalle"],
        "Whisper": ["whisper", "openai whisper"],
        "LLaMA": ["llama", "llama-2", "llama-3", "llama 2", "llama 3"],
        "Mistral": ["mistral"],
        "GPT": ["gpt-4", "gpt-3", "gpt-4o", "gpt-3.5-turbo"],
        "BERT": ["bert", "roberta", "distilbert"],
        "Word2Vec": ["word2vec"],
        "GloVe": ["glove"],
        "FastText": ["fasttext"],
        "TF-IDF": ["tf-idf", "tfidf"],
        "Pandas": ["pandas"],
        "NumPy": ["numpy"],
        "SciPy": ["scipy"],
        "Matplotlib": ["matplotlib"],
        "Seaborn": ["seaborn"],
        "Plotly": ["plotly"],
        "MLflow": ["mlflow"],
        "Weights & Biases": ["wandb", "weights and biases", "weights & biases"],
        "Ray": ["ray", "ray tune"],
        "Dask": ["dask"],
        "Apache Spark": ["spark", "pyspark", "apache spark"],
        "Apache Airflow": ["airflow", "apache airflow"],
        "Kubeflow": ["kubeflow"],
        "Vertex AI": ["vertex ai"],
        "SageMaker": ["sagemaker", "aws sagemaker"],
        "Azure ML": ["azure ml", "azure machine learning"],
        "AutoML": ["automl", "auto ml"],
        "Reinforcement Learning": ["reinforcement learning", "rl"],
        "GANs": ["gan", "gans", "generative adversarial"],
        "RAG": ["rag", "retrieval augmented generation"],
        "Vector Embeddings": ["vector embeddings", "embeddings"],
    },
}


# =============================================================================
# LAYER 3a: AHO-CORASICK AUTOMATON — Multi-pattern string matching
# =============================================================================

class AhoCorasickAutomaton:
    """
    Aho-Corasick algorithm implementation for efficient multi-pattern matching.
    Builds a finite automaton from all keyword patterns and scans text in O(n + m + z)
    where n = text length, m = total pattern chars, z = number of matches.
    """

    def __init__(self):
        self.goto_fn = [{}]          # goto function (trie transitions)
        self.failure_fn = [0]        # failure function (suffix links)
        self.output_fn = defaultdict(set)  # output function (matched patterns at each state)
        self.num_states = 1

    def _add_pattern(self, pattern, label):
        """Add a single pattern to the automaton's trie."""
        state = 0
        for char in pattern.lower():
            if char not in self.goto_fn[state]:
                self.goto_fn[state][char] = self.num_states
                self.goto_fn.append({})
                self.failure_fn.append(0)
                self.num_states += 1
            state = self.goto_fn[state][char]
        self.output_fn[state].add(label)

    def build(self, patterns_with_labels):
        """
        Build the automaton from a list of (pattern, label) tuples.
        Computes failure links using BFS (Breadth-First Search).
        """
        for pattern, label in patterns_with_labels:
            self._add_pattern(pattern, label)

        queue = deque()

        for char, next_state in self.goto_fn[0].items():
            self.failure_fn[next_state] = 0
            queue.append(next_state)

        while queue:
            current_state = queue.popleft()
            for char, next_state in self.goto_fn[current_state].items():
                queue.append(next_state)
                failure_state = self.failure_fn[current_state]
                while failure_state != 0 and char not in self.goto_fn[failure_state]:
                    failure_state = self.failure_fn[failure_state]
                self.failure_fn[next_state] = self.goto_fn[failure_state].get(char, 0)
                if self.failure_fn[next_state] != next_state:
                    self.output_fn[next_state] |= self.output_fn[self.failure_fn[next_state]]

    def search(self, text):
        """
        Search text for all patterns in a single pass.
        Ensures word boundaries to avoid false positives (e.g., 'go' inside 'algorithm').
        """
        state = 0
        matches = set()
        text_lower = text.lower()
        n = len(text_lower)

        for i, char in enumerate(text_lower):
            while state != 0 and char not in self.goto_fn[state]:
                state = self.failure_fn[state]
            state = self.goto_fn[state].get(char, 0)

            if self.output_fn[state]:
                for label, pat_len in self.output_fn[state]:
                    start_pos = i - pat_len + 1
                    end_pos = i + 1
                    # Word boundary check
                    prev_char = text_lower[start_pos - 1] if start_pos > 0 else ' '
                    next_char = text_lower[end_pos] if end_pos < n else ' '
                    
                    is_prev_boundary = not (prev_char.isalnum() or prev_char in ['+', '#', '.'])
                    is_next_boundary = not (next_char.isalnum() or next_char in ['+', '#'])

                    if is_prev_boundary and is_next_boundary:
                        matches.add(label)

        return matches


def build_tech_automaton():
    """
    Build Aho-Corasick automaton from the entire TECH_DICTIONARY.
    Returns (automaton, alias_to_tech_map).
    """
    automaton = AhoCorasickAutomaton()
    alias_to_tech = {}
    patterns = []

    for category, techs in TECH_DICTIONARY.items():
        for tech_name, aliases in techs.items():
            label = f"{category}::{tech_name}"
            for alias in aliases:
                patterns.append((alias, (label, len(alias))))
                alias_to_tech[label] = (category, tech_name)

    automaton.build(patterns)
    return automaton, alias_to_tech


# Build the global automaton once at module load time
TECH_AUTOMATON, ALIAS_TO_TECH = build_tech_automaton()


# =============================================================================
# LAYER 3b: TF-IDF PARAGRAPH SCORER — Ranks paragraphs by keyword relevance
# =============================================================================

class TFIDFParagraphScorer:
    """
    TF-IDF (Term Frequency - Inverse Document Frequency) scorer for
    ranking paragraphs by relevance to a set of target keywords.
    """

    def __init__(self, paragraphs):
        self.paragraphs = paragraphs
        self.N = len(paragraphs)
        self.tokenized = [self._tokenize(p) for p in paragraphs]
        self.doc_freq = self._compute_doc_freq()

    def _tokenize(self, text):
        return re.findall(r'\b[a-z]{2,}\b', text.lower())

    def _compute_doc_freq(self):
        df = defaultdict(int)
        for tokens in self.tokenized:
            unique_tokens = set(tokens)
            for token in unique_tokens:
                df[token] += 1
        return df

    def _tf(self, term, doc_tokens):
        if not doc_tokens:
            return 0.0
        return doc_tokens.count(term) / len(doc_tokens)

    def _idf(self, term):
        df = self.doc_freq.get(term, 0)
        return math.log((self.N + 1) / (1 + df))

    def score_paragraphs(self, target_keywords):
        scored = []
        keywords_lower = [kw.lower() for kw in target_keywords]

        for idx, tokens in enumerate(self.tokenized):
            score = 0.0
            for keyword in keywords_lower:
                tf = self._tf(keyword, tokens)
                idf = self._idf(keyword)
                score += tf * idf
            if score > 0:
                scored.append((idx, score))

        scored.sort(key=lambda x: x[1], reverse=True)
        return scored


# =============================================================================
# LAYER 1: TEXT EXTRACTION — File → Raw Text
# =============================================================================

def extract_text_from_pdf(file_path):
    """Extract text from PDF using pdfplumber (layout-aware)."""
    import pdfplumber
    text_parts = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {e}")
    return '\n\n'.join(text_parts)


def extract_text_from_docx(file_path):
    """Extract text from DOCX using python-docx with style awareness."""
    from docx import Document
    doc = Document(file_path)
    text_parts = []
    structured_parts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style else 'Normal'
        text_parts.append(text)
        structured_parts.append({
            'text': text,
            'style': style_name,
            'is_heading': style_name.lower().startswith('heading') or style_name.lower() == 'title',
        })

    return '\n'.join(text_parts), structured_parts


def extract_text_from_txt(file_path):
    """Extract text from plain text file with encoding fallback."""
    encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']
    for enc in encodings:
        try:
            with open(file_path, 'r', encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError("Could not decode text file with any supported encoding.")


def extract_text(file_path):
    """Main text extraction dispatcher."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        text = extract_text_from_pdf(file_path)
        return text, None
    elif ext in ('.docx', '.doc'):
        return extract_text_from_docx(file_path)
    elif ext in ('.txt', '.md', '.rst'):
        text = extract_text_from_txt(file_path)
        return text, None
    else:
        raise ValueError(f"Unsupported file format: {ext}. Supported: PDF, DOCX, TXT")


# =============================================================================
# LAYER 2: SECTION-AWARE REGEX PARSER — Structured field extraction
# =============================================================================

SECTION_PATTERNS = {
    'title': re.compile(
        r'(?:project\s*title|title|project\s*name|proposal\s*title|name\s*of\s*(?:the\s*)?project)'
        r'\s*[:：\-–—]\s*([\s\S]+?)(?=\n\s*(?:[A-Z][a-z0-9\s]{1,30}[:：\-–—]|\n\n|\Z))',
        re.IGNORECASE
    ),
    'description': re.compile(
        r'(?:description|abstract|executive\s*summary|project\s*summary|overview|introduction|synopsis)'
        r'\s*[:：\-–—]\s*([\s\S]+?)(?=\n\s*(?:[A-Z][a-z0-9\s]{2,25}[:：\-–—]|\n\n|\Z))',
        re.IGNORECASE
    ),
    'problem_statement': re.compile(
        r'(?:problem\s*statement|problem\s*definition|problem|research\s*problem|challenges?|background\s*(?:&|and)\s*problem)'
        r'\s*[:：\-–—]\s*([\s\S]+?)(?=\n\s*(?:[A-Z][a-z0-9\s]{2,25}[:：\-–—]|\n\n|\Z))',
        re.IGNORECASE
    ),
    'objectives': re.compile(
        r'(?:objectives?|goals?\s*(?:&|and)\s*objectives?|project\s*goals?|key\s*objectives?|aims?|'
        r'deliverables?|milestones?|expected\s*outcomes?|scope)'
        r'\s*[:：\-–—]\s*([\s\S]+?)(?=\n\s*(?:[A-Z][a-z0-9\s]{2,25}[:：\-–—]|\n\n|\Z))',
        re.IGNORECASE
    ),
    'tech_stack_raw': re.compile(
        r'(?:tech(?:nology|nical)?\s*stack|technologies?\s*used|tools?\s*(?:&|and)\s*technologies|'
        r'system\s*requirements|technical\s*specifications?|tech\s*used|software\s*used|'
        r'programming\s*(?:languages?|tools?))'
        r'\s*[:：\-–—]\s*([\s\S]+?)(?=\n\s*(?:[A-Z][a-z0-9\s]{2,25}[:：\-–—]|\n\n|\Z))',
        re.IGNORECASE
    ),
}

GITHUB_URL_PATTERN = re.compile(
    r'https?://(?:www\.)?github\.com/[\w\-\.]+/[\w\-\.]+(?:/[\w\-\./]*)?',
    re.IGNORECASE
)
GITLAB_URL_PATTERN = re.compile(
    r'https?://(?:www\.)?gitlab\.com/[\w\-\.]+/[\w\-\.]+(?:/[\w\-\./]*)?',
    re.IGNORECASE
)
DOC_URL_PATTERN = re.compile(
    r'https?://(?:docs|confluence|notion|wiki|readme|gitbook|readthedocs)[\w\.\-/]+',
    re.IGNORECASE
)
BOILERPLATE_PATTERNS = re.compile(
    r'(?:bonafide\s*certificate|certificate|acknowledgement|anna\s*university|'
    r'bachelor\s*of\s*engineering|department\s*of\s*computer\s*science|'
    r'in\s*partial\s*fulfillment|submitted\s*by|internal\s*examiner|'
    r'external\s*examiner|viva-voce|table\s*of\s*contents|list\s*of\s*figures|'
    r'list\s*of\s*tables|declaration|head\s*of\s*the\s*department|internal\s*guide)',
    re.IGNORECASE
)


def is_boilerplate(text):
    """Check if text is boilerplate front-matter (certificates, acknowledgements)."""
    if not text:
        return True
    return bool(BOILERPLATE_PATTERNS.search(text[:350]))


def extract_heading_section(raw_text, heading_patterns, next_patterns):
    """
    Extracts text under a standalone or numbered heading:
    e.g. ABSTRACT\n[Content] or 3.1 Problem Statement\n[Content]
    """
    for hp in heading_patterns:
        regex = re.compile(
            rf'(?:^|\n)\s*{hp}\s*[:：\-–—]?\s*\n+([\s\S]+?)(?=(?:\n\s*(?:{"|".join(next_patterns)})\b)|\Z)',
            re.IGNORECASE | re.MULTILINE
        )
        match = regex.search(raw_text)
        if match:
            val = match.group(1).strip()
            # Remove isolated page number lines
            val = re.sub(r'\n\s*\d+\s*\n', '\n', val)
            val = re.sub(r'\s+', ' ', val).strip()
            if len(val) > 20 and not is_boilerplate(val):
                return val
    return None


def extract_structured_fields(raw_text):
    """
    Layer 2: Scan for explicitly labeled sections and standalone heading sections.
    """
    fields = {}
    confidence = {}

    # 1. Try labeled sections (colon-delimited)
    for field_name, pattern in SECTION_PATTERNS.items():
        match = pattern.search(raw_text)
        if match:
            value = match.group(1).strip()
            value = re.sub(r'\s+', ' ', value).strip()
            if len(value) > 10 and not is_boilerplate(value[:250]):
                fields[field_name] = value
                confidence[field_name] = 'high'

    # 2. Try standalone heading sections (academic project reports, PRDs, theses)
    if 'description' not in fields:
        abstract_val = extract_heading_section(
            raw_text,
            [r'abstract', r'executive\s*summary', r'project\s*summary', r'overview', r'synopsis'],
            [r'keywords', r'chapter\s*\d+', r'table\s*of\s*contents', r'1\.\s*introduction', r'introduction']
        )
        if abstract_val:
            fields['description'] = abstract_val
            confidence['description'] = 'high'

    if 'problem_statement' not in fields:
        problem_val = extract_heading_section(
            raw_text,
            [r'(?:\d+\.\d+\s*)?problem\s*statement', r'(?:\d+\.\d+\s*)?problem\s*definition', r'problem\s*description'],
            [r'(?:\d+\.\d+\s*)?objectives?', r'chapter\s*\d+', r'\d+\.\d+\s+[A-Z]']
        )
        if problem_val:
            fields['problem_statement'] = problem_val
            confidence['problem_statement'] = 'high'

    if 'objectives' not in fields:
        obj_val = extract_heading_section(
            raw_text,
            [r'(?:\d+\.\d+\s*)?objectives?', r'(?:\d+\.\d+\s*)?project\s*objectives?', r'(?:\d+\.\d+\s*)?goals?'],
            [r'chapter\s*\d+', r'\d+\.\d+\s+[A-Z]', r'existing\s*system', r'scope']
        )
        if obj_val:
            fields['objectives'] = obj_val
            confidence['objectives'] = 'high'

    github_match = GITHUB_URL_PATTERN.search(raw_text)
    if github_match:
        fields['github_url'] = github_match.group(0)
        confidence['github_url'] = 'high'

    if not fields.get('github_url'):
        gitlab_match = GITLAB_URL_PATTERN.search(raw_text)
        if gitlab_match:
            fields['github_url'] = gitlab_match.group(0)
            confidence['github_url'] = 'high'

    doc_match = DOC_URL_PATTERN.search(raw_text)
    if doc_match:
        fields['documentation_url'] = doc_match.group(0)
        confidence['documentation_url'] = 'high'

    return fields, confidence


# =============================================================================
# LAYER 3: HEURISTIC KEYWORD MATCHER — Fallback for missing fields
# =============================================================================

PROBLEM_KEYWORDS = [
    'problem', 'challenge', 'bottleneck', 'issue', 'limitation',
    'gap', 'inefficiency', 'difficulty', 'obstacle', 'constraint',
    'shortcoming', 'deficiency', 'weakness', 'drawback', 'barrier',
    'pain', 'manual', 'error', 'slow', 'outdated', 'lack',
    'insufficient', 'unreliable', 'complex', 'costly'
]

OBJECTIVE_KEYWORDS = [
    'objective', 'goal', 'aim', 'target', 'milestone',
    'deliverable', 'outcome', 'purpose', 'achieve', 'develop',
    'implement', 'design', 'build', 'create', 'improve',
    'reduce', 'increase', 'optimize', 'automate', 'enhance',
    'provide', 'enable', 'ensure', 'establish', 'integrate'
]


def split_into_paragraphs(text):
    """
    Split text into paragraphs. Handles both double newlines (\n\n)
    and single newline documents where paragraphs are separate lines.
    """
    paragraphs = [p.strip() for p in re.split(r'\n\s*\n', text) if p.strip() and len(p.strip()) > 20]
    if len(paragraphs) <= 1:
        lines = [line.strip() for line in text.split('\n') if len(line.strip()) > 25]
        if len(lines) > 1:
            paragraphs = lines
    return paragraphs


def infer_title(raw_text, structured_parts=None):
    # Check for "report titled [title]" or "project titled [title]"
    quoted_match = re.search(r'(?:report|project)\s*titled\s*["“](.+?)["”]', raw_text, re.IGNORECASE)
    if quoted_match:
        cand = quoted_match.group(1).strip()
        if 10 < len(cand) < 180 and not is_boilerplate(cand):
            return cand.title() if cand.isupper() else cand, 'high'

    if structured_parts:
        for part in structured_parts:
            if part['is_heading'] and 10 < len(part['text']) < 150 and not is_boilerplate(part['text']):
                return part['text'], 'medium'

    lines = [line.strip() for line in raw_text.strip().split('\n') if line.strip()]
    title_candidates = []
    for line in lines[:8]:
        if re.search(r'(?:phase|submitted\s*by|bachelor\s*of|anna\s*university|bonafide|author\s*:|date\s*:|page\s*\d+)', line, re.IGNORECASE):
            break
        if 5 < len(line) < 120 and not is_boilerplate(line):
            title_candidates.append(line)
            if len(' '.join(title_candidates)) > 55:
                break
    
    if title_candidates:
        full_title = ' '.join(title_candidates).strip()
        full_title = re.sub(r'\s+', ' ', full_title)
        if 10 < len(full_title) < 200 and not is_boilerplate(full_title):
            return full_title, 'medium'

    return None, None


def infer_description(raw_text, structured_parts=None, title=None):
    paragraphs = split_into_paragraphs(raw_text)
    for para in paragraphs:
        if len(para) > 60:
            if is_boilerplate(para):
                continue
            if title and (para.strip().startswith(title[:25]) or para.strip() == title.strip()):
                continue
            return para[:1200], 'medium'
    return None, None


def infer_problem_statement(raw_text, title=None):
    paragraphs = split_into_paragraphs(raw_text)
    if not paragraphs:
        return None, None

    # Filter out title and boilerplate
    candidates = [p for p in paragraphs if not is_boilerplate(p) and not (title and (p == title or p.startswith(title[:25])))]
    if not candidates:
        candidates = [p for p in paragraphs if not is_boilerplate(p)]
    if not candidates:
        return None, None

    scorer = TFIDFParagraphScorer(candidates)
    scored = scorer.score_paragraphs(PROBLEM_KEYWORDS)

    if scored and scored[0][1] > 0.0005:
        best_idx = scored[0][0]
        best_score = scored[0][1]
        confidence = 'medium' if best_score > 0.02 else 'low'
        return candidates[best_idx][:1000], confidence

    return None, None


def infer_objectives(raw_text, title=None):
    paragraphs = split_into_paragraphs(raw_text)
    if not paragraphs:
        return None, None

    candidates = [p for p in paragraphs if not is_boilerplate(p) and not (title and (p == title or p.startswith(title[:25])))]
    if not candidates:
        candidates = [p for p in paragraphs if not is_boilerplate(p)]
    if not candidates:
        return None, None

    list_pattern = re.compile(r'(?:^\s*(?:\d+[\.\)]\s+|[•\-\*]\s+).+\n?)+', re.MULTILINE)
    list_paragraphs = []
    for i, para in enumerate(candidates):
        if list_pattern.search(para):
            list_paragraphs.append((i, para))

    scorer = TFIDFParagraphScorer(candidates)
    scored = scorer.score_paragraphs(OBJECTIVE_KEYWORDS)

    list_indices = {idx for idx, _ in list_paragraphs}
    boosted = []
    for idx, score in scored:
        if idx in list_indices:
            score *= 1.5
        boosted.append((idx, score))
    boosted.sort(key=lambda x: x[1], reverse=True)

    if boosted and boosted[0][1] > 0.0005:
        best_idx = boosted[0][0]
        best_score = boosted[0][1]
        confidence = 'medium' if best_score > 0.02 else 'low'
        return candidates[best_idx][:1500], confidence

    return None, None


def detect_technologies(raw_text):
    matches = TECH_AUTOMATON.search(raw_text)

    result = {
        'programming_languages': [],
        'frameworks': [],
        'database_tech': [],
        'apis_used': [],
        'ai_ml_tech': [],
    }

    seen = set()
    for label in matches:
        category, tech_name = ALIAS_TO_TECH[label]
        if tech_name not in seen:
            result[category].append(tech_name)
            seen.add(tech_name)

    return result


# =============================================================================
# MAIN PIPELINE — Orchestrates all 3 layers
# =============================================================================

def extract_from_document(file_path):
    """
    Main extraction pipeline. Runs all 3 layers and produces
    a structured result matching the Project model fields.
    """
    raw_text, structured_parts = extract_text(file_path)

    if not raw_text or len(raw_text.strip()) < 20:
        return {
            'extracted_fields': {},
            'confidence': {},
            'raw_text_preview': raw_text[:500] if raw_text else '',
            'extraction_summary': 'Document contains insufficient text for extraction.',
        }

    fields, confidence = extract_structured_fields(raw_text)

    if 'title' not in fields:
        title, conf = infer_title(raw_text, structured_parts)
        if title:
            fields['title'] = title
            confidence['title'] = conf

    if 'description' not in fields:
        desc, conf = infer_description(raw_text, structured_parts, fields.get('title'))
        if desc:
            fields['description'] = desc
            confidence['description'] = conf

    if 'problem_statement' not in fields:
        problem, conf = infer_problem_statement(raw_text, fields.get('title'))
        if problem:
            fields['problem_statement'] = problem
            confidence['problem_statement'] = conf

    if 'objectives' not in fields:
        objectives, conf = infer_objectives(raw_text, fields.get('title'))
        if objectives:
            fields['objectives'] = objectives
            confidence['objectives'] = conf

    tech_results = detect_technologies(raw_text)

    if 'tech_stack_raw' in fields:
        extra_techs = detect_technologies(fields['tech_stack_raw'])
        for category in tech_results:
            for tech in extra_techs.get(category, []):
                if tech not in tech_results[category]:
                    tech_results[category].append(tech)
        del fields['tech_stack_raw']

    for category, techs in tech_results.items():
        if techs:
            fields[category] = techs
            confidence[category] = 'high'

    total_fields = len([v for v in fields.values() if v])
    high_conf = len([v for v in confidence.values() if v == 'high'])
    med_conf = len([v for v in confidence.values() if v == 'medium'])
    low_conf = len([v for v in confidence.values() if v == 'low'])

    summary_parts = [f"Extracted {total_fields} fields"]
    if high_conf:
        summary_parts.append(f"{high_conf} high confidence")
    if med_conf:
        summary_parts.append(f"{med_conf} medium confidence")
    if low_conf:
        summary_parts.append(f"{low_conf} low confidence")

    return {
        'extracted_fields': fields,
        'confidence': confidence,
        'raw_text_preview': raw_text[:500],
        'extraction_summary': ' — '.join(summary_parts),
    }
